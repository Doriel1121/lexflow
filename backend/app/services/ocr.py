from typing import Optional
import os
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

class OCRService:
    def __init__(self):
        pass

    def _looks_like_corrupted_pdf_text(self, text: str, page_count: int) -> bool:
        """Detect PDFs that expose long but unusable embedded text.

        Some RTL/Hebrew PDFs contain extractable text streams, but the glyph mapping is
        broken/mirrored and full of repeated fragments. Length alone is therefore not
        a safe signal that OCR can be skipped.
        """
        clean = text or ""
        stripped = clean.strip()
        if not stripped:
            return True

        # If the PDF extracted enough text and it contains Hebrew plus obvious
        # non-Hebrew glyph-map artifacts, prefer OCR. This catches the LegalOS
        # lease sample without forcing OCR for normal English PDFs.
        sample = stripped[: min(len(stripped), 8000)]
        total_chars = max(1, len(sample))
        hebrew_chars = sum(1 for ch in sample if "\u0590" <= ch <= "\u05FF")
        armenian_artifacts = sum(1 for ch in sample if "\u0530" <= ch <= "\u058F")
        control_artifacts = sum(1 for ch in sample if ord(ch) < 32 and ch not in "\n\r\t\f")

        repeated_phrases = 0
        words = [word for word in sample.split() if len(word) >= 5]
        if words:
            from collections import Counter
            counts = Counter(words)
            repeated_phrases = sum(1 for _, count in counts.most_common(12) if count >= max(4, page_count + 1))

        artifact_ratio = (armenian_artifacts + control_artifacts) / total_chars
        has_hebrew = hebrew_chars >= 20
        return has_hebrew and (artifact_ratio >= 0.01 or repeated_phrases >= 3)
    
    async def extract_text_from_file(self, file_path: str) -> dict:
        """Extract text from file with metadata.
        
        For B2/R2 storage: Automatically downloads file to temp location before processing.
        """
        from app.services.file_processor import FileProcessor
        
        logger.info(f"Starting OCR extraction for: {file_path}")
        
        # If using B2/R2, download file to temp location first
        # Use async version since we're in an async context
        actual_file_path = await FileProcessor.get_processing_file_path_async(file_path)
        is_temp_cloud_file = (actual_file_path != file_path)
        
        logger.info(f"Processing file path: {actual_file_path} (temp={is_temp_cloud_file})")
        
        try:
            path = Path(actual_file_path)
            if not path.exists():
                logger.error(f"FILE NOT FOUND: {actual_file_path} (original: {file_path})")
                # Don't return placeholder - return explicit error for debugging
                return {
                    "text": "",
                    "language": "en",
                    "page_count": 0,
                    "error": f"File not found: {actual_file_path}"
                }
            
            # Text files
            if path.suffix.lower() in ['.txt', '.md', '.csv', '.log']:
                try:
                    with open(path, 'r', encoding='utf-8') as f:
                        text = f.read()
                        logger.info(f"Successfully extracted text from {path.suffix} file: {len(text)} characters")
                        return {"text": text, "language": "en", "page_count": 1}
                except UnicodeDecodeError as e:
                    logger.warning(f"Text file encoding error for {path.name}: {e}. Trying latin-1 encoding...")
                    try:
                        with open(path, 'r', encoding='latin-1') as f:
                            text = f.read()
                            logger.info(f"Successfully extracted text with latin-1: {len(text)} characters")
                            return {"text": text, "language": "en", "page_count": 1}
                    except Exception as e2:
                        logger.error(f"Failed to read text file with latin-1: {e2}", exc_info=True)
                except Exception as e:
                    logger.error(f"Failed to extract text from {path.name}: {type(e).__name__}: {e}", exc_info=True)
            
            # PDF files
            if path.suffix.lower() == '.pdf':
                try:
                    import PyPDF2
                    text_parts = []
                    page_count = 0
                    with open(path, 'rb') as f:
                        reader = PyPDF2.PdfReader(f)
                        page_count = len(reader.pages)
                        for idx, page in enumerate(reader.pages):
                            page_text = page.extract_text() or ""
                            text_parts.append(page_text)
                    # Join pages with an explicit page break marker so downstream
                    # services (chunker, normalization) can preserve page numbers.
                    joined_text = "\n\n--- Page Break ---\n\n".join(text_parts)

                    # Hybrid Routing Logic
                    # If this is a digital PDF, text length will be substantial. 
                    # If it's a scanned PDF, PyPDF2 extracts virtually nothing (or just noise).
                    # We expect at least ~100 characters per page for a real digital document.
                    text_length = len(joined_text.strip())
                    threshold = max(200, page_count * 100)
                    
                    if text_length < threshold or self._looks_like_corrupted_pdf_text(joined_text, page_count):
                        logger.info(
                            "PDF text extraction looked weak/corrupted (chars=%s, threshold=%s, pages=%s). Falling back to OCR.",
                            text_length,
                            threshold,
                            page_count,
                        )
                        from app.services.ocr_engine import tesseract_ocr_service
                        return await tesseract_ocr_service.extract_text_from_scanned_pdf(actual_file_path)
                    
                    return {"text": joined_text, "language": "en", "page_count": page_count}
                except Exception as e:
                    # If PyPDF2 crashes entirely (corrupted metadata etc), fallback to OCR
                    from app.services.ocr_engine import tesseract_ocr_service
                    return await tesseract_ocr_service.extract_text_from_scanned_pdf(actual_file_path)
            
            # DOCX/DOC files
            if path.suffix.lower() in ['.docx', '.doc']:
                try:
                    from docx import Document
                    doc = Document(path)
                    text_parts = [para.text for para in doc.paragraphs]
                    # Also extract text from tables
                    for table in doc.tables:
                        for row in table.rows:
                            for cell in row.cells:
                                text_parts.append(cell.text)
                    full_text = "\n".join(text_parts)
                    logger.info(f"Successfully extracted text from DOCX: {len(full_text)} characters")
                    return {"text": full_text, "language": "en", "page_count": 1}
                except Exception as e:
                    logger.error(f"Failed to extract text from DOCX: {e}", exc_info=True)
                    return {"text": "", "language": "en", "page_count": 0}
            
            # Other files
            return {"text": f"Document: {path.name}\n\nFile Type: {path.suffix}\nSize: {path.stat().st_size} bytes", "language": "en", "page_count": 1}
        except Exception as e:
            return {"text": f"Document uploaded: {path.name if 'path' in locals() else 'file'}", "language": "en", "page_count": 0}
        finally:
            # Clean up B2/R2 temp file if downloaded
            if is_temp_cloud_file:
                FileProcessor.cleanup_temp_file(actual_file_path)
    
    async def extract_text(self, file_path: str, language: Optional[str] = None) -> dict:
        return await self.extract_text_from_file(file_path)

    async def detect_language(self, file_path: str) -> str:
        return "en"

ocr_service = OCRService()
